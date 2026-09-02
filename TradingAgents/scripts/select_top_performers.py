"""Select the top 20% of tickers per sector from the AI-readiness index workbook.

Ranks by the workbook's ``Index`` column (the only per-ticker scoring column
present) within each ``Sector`` group, and writes the selection to CSV for the
batch runner to consume.

Usage:
    python scripts/select_top_performers.py \\
        --input data/aide_index.xlsx \\
        --output output/top_performers.csv \\
        --pct 0.20
"""

import argparse
import csv
import math
from collections import defaultdict
from pathlib import Path

import openpyxl

REQUIRED_COLUMNS = ("Sector", "Ticker", "Industry", "Index")


def load_rows(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {name: header.index(name) for name in REQUIRED_COLUMNS if name in header}
    missing = [name for name in REQUIRED_COLUMNS if name not in col_idx]
    if missing:
        raise ValueError(f"Workbook is missing required column(s): {missing}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        sector = r[col_idx["Sector"]]
        ticker = r[col_idx["Ticker"]]
        industry = r[col_idx["Industry"]]
        index = r[col_idx["Index"]]
        if not sector or not ticker or index is None:
            continue  # skip blank trailer rows
        rows.append({"Sector": sector, "Ticker": ticker, "Industry": industry, "Index": float(index)})
    return rows


def select_top_pct(rows: list[dict], pct: float) -> list[dict]:
    by_sector = defaultdict(list)
    for row in rows:
        by_sector[row["Sector"]].append(row)

    selected = []
    for sector, items in sorted(by_sector.items()):
        items.sort(key=lambda r: r["Index"], reverse=True)
        k = max(1, math.ceil(len(items) * pct))
        for rank, row in enumerate(items[:k], start=1):
            selected.append({**row, "SectorRank": rank, "SectorCount": len(items)})
    return selected


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="data/aide_index.xlsx", help="Source workbook")
    parser.add_argument("--output", default="output/top_performers.csv", help="Destination CSV")
    parser.add_argument("--pct", type=float, default=0.20, help="Top fraction per sector (default 0.20)")
    args = parser.parse_args()

    rows = load_rows(Path(args.input))
    selected = select_top_pct(rows, args.pct)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["Sector", "Ticker", "Industry", "Index", "SectorRank", "SectorCount"])
        writer.writeheader()
        writer.writerows(selected)

    print(f"Loaded {len(rows)} tickers across {len({r['Sector'] for r in rows})} sectors.")
    print(f"Selected {len(selected)} tickers (top {args.pct:.0%} per sector) -> {out_path}")
    by_sector_count = defaultdict(int)
    for r in selected:
        by_sector_count[r["Sector"]] += 1
    for sector, count in sorted(by_sector_count.items()):
        print(f"  {sector}: {count}")


if __name__ == "__main__":
    main()
